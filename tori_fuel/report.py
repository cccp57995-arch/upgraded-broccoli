"""加油計畫 Excel 報表（5 頁）。

1. 用油分析摘要 — 判定規則、歷史統計、各類型油耗率、建議採用值
2. 月份統計     — 歷史實際月份用油（出航/靠港/進塢分色）
3. 每日明細     — 歷史逐日油耗（含加油量欄，加油日黃底）
4. 2026加油預估 — 依航次類型推算油耗、加油時機與加油量
5. 2026逐日ROB  — 預估 ROB、實際 ROB（有資料才填）、日耗油誤差（實際−預估）
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import voyage_types
from .history import DailyRecord, REFUEL_DELTA_KL, STATUS_DOCK, STATUS_PORT, STATUS_SAIL
from .plan2026 import Plan2026, TANK_CAPACITY_KL, WARNING_LEVEL_KL
from .typestats import RateTable, monthly_summary

FILL_HEADER = PatternFill("solid", fgColor="0B3D62")
FILL_TITLE = PatternFill("solid", fgColor="D7E0E8")
FILL_SAIL = PatternFill("solid", fgColor="DCE9F7")   # 出航 藍
FILL_PORT = PatternFill("solid", fgColor="E2F0DC")   # 靠港 綠
FILL_DOCK = PatternFill("solid", fgColor="E4E4E4")   # 進塢 灰
FILL_REFUEL = PatternFill("solid", fgColor="FFF2A8")  # 加油 黃
FILL_ACTUAL = PatternFill("solid", fgColor="F2F2F2")  # 已有實際資料 淡灰
FONT_HEADER = Font(bold=True, color="FFFFFF")
FONT_TITLE = Font(bold=True, size=12)
FONT_ALERT = Font(bold=True, color="C00000")
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

STATUS_FILL = {STATUS_SAIL: FILL_SAIL, STATUS_PORT: FILL_PORT, STATUS_DOCK: FILL_DOCK}


def _header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=col, value=label)
        c.fill, c.font, c.border, c.alignment = FILL_HEADER, FONT_HEADER, THIN, CENTER


def _title(ws, row: int, text: str, span: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill = FONT_TITLE, FILL_TITLE


def _widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _sheet_summary(wb: Workbook, rates: RateTable,
                   records: list[DailyRecord]) -> None:
    ws = wb.create_sheet("用油分析摘要")
    _widths(ws, [26, 14, 12, 12, 12, 12, 12, 42])
    r = 1
    _title(ws, r, "TORI 勵進研究船 2026 加油計畫 — 用油分析摘要", 8); r += 2

    ws.cell(row=r, column=1, value="一、判定規則").font = FONT_TITLE; r += 1
    rules = [
        "航次日分類（順序固定）：① 當日總油耗 = 0 → 進塢；② 推進時數 > 0 或 作業時數 > 0 "
        "或 距離 ≠ 0 或 推進油耗 > 0 或 作業油耗 > 0 或 靠港油耗 = 0 → 出航；③ 其餘 → 靠港",
        f"加油判定：前後兩日 ROB 差值 > {REFUEL_DELTA_KL:.0f} KL 即視為加油，自動計算加油量",
        f"ROB 參數：油艙上限 {TANK_CAPACITY_KL:.0f} KL／警戒線 {WARNING_LEVEL_KL:.0f} KL／"
        "加油策略補滿至上限（已發生加油採實際加油量）",
        "加油時機：航次備註含「出塢加油」→ 塢修結束次日；含「加燃油」→ 該航次返港次日",
    ]
    for text in rules:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.cell(row=r, column=1, value="• " + text).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="二、歷史全體統計").font = FONT_TITLE; r += 1
    first, last = records[0].day, records[-1].day
    ws.cell(row=r, column=1,
            value=f"資料期間：{first:%Y/%m/%d} – {last:%Y/%m/%d}（共 {len(records)} 天）")
    r += 1
    _header(ws, r, ["分類", "天數", "日均(KL)", "P25", "P50", "P75", "", ""]); r += 1
    ov = rates.overall
    for label, fuels, fill in [
        (STATUS_SAIL, ov.sail_fuels, FILL_SAIL),
        (STATUS_PORT, ov.port_fuels, FILL_PORT),
    ]:
        from .typestats import percentile
        mean = sum(fuels) / len(fuels) if fuels else 0.0
        vals = [label, len(fuels), round(mean, 2),
                round(percentile(fuels, 25), 2), round(percentile(fuels, 50), 2),
                round(percentile(fuels, 75), 2), "", ""]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = THIN
            if col == 1:
                c.fill = fill
        r += 1
    dock_days = sum(1 for rec in records if rec.status == STATUS_DOCK)
    for col, v in enumerate([STATUS_DOCK, dock_days, 0, 0, 0, 0, "", ""], 1):
        c = ws.cell(row=r, column=col, value=v)
        c.border = THIN
        if col == 1:
            c.fill = FILL_DOCK
    r += 2

    ws.cell(row=r, column=1, value="三、各任務類型出航日油耗率").font = FONT_TITLE; r += 1
    _header(ws, r, ["任務類型", "出航日數", "日均(KL)", "P25", "P50", "P75",
                    "建議採用值", "採用依據"]); r += 1
    for vtype in voyage_types.ALL_TYPES:
        st = rates.by_type.get(vtype)
        if st and st.sail_days:
            vals = [vtype, st.sail_days, round(st.sail_mean, 2),
                    round(st.sail_p(25), 2), round(st.sail_p(50), 2),
                    round(st.sail_p(75), 2),
                    rates.sail_rates[vtype], rates.sail_basis[vtype]]
        else:
            vals = [vtype, 0, "—", "—", "—", "—",
                    rates.sail_rates[vtype], rates.sail_basis[vtype]]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = THIN
            if col == 7:
                c.font = Font(bold=True)
                c.fill = FILL_REFUEL
        r += 1
    for col, v in enumerate(["靠港日（全類型）", len(ov.port_fuels),
                             round(ov.port_mean, 2), "", "", "",
                             rates.port_rate, "全體靠港日中位數"], 1):
        c = ws.cell(row=r, column=col, value=v)
        c.border = THIN
        if col == 7:
            c.font = Font(bold=True)
            c.fill = FILL_REFUEL
    r += 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.cell(row=r, column=1, value=(
        "註：歷史航次中僅 2026 年航次與 T 航次（試航）可由航次表確認任務類型，"
        "2024–2025 航次列入全體統計。震測等無分類樣本之類型暫採全體出航日中位數，"
        "待補歷史航次總表後自動改採類型實績。")).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r].height = 42


def _sheet_monthly(wb: Workbook, records: list[DailyRecord]) -> None:
    ws = wb.create_sheet("月份統計")
    _widths(ws, [10, 9, 9, 9, 12, 12, 12, 11])
    _title(ws, 1, "歷史月份用油統計（出航＝藍、靠港＝綠、進塢＝灰）", 8)
    _header(ws, 2, ["月份", "出航天數", "靠港天數", "進塢天數",
                    "出航用油(KL)", "靠港用油(KL)", "總用油(KL)", "加油量(KL)"])
    r = 3
    for m in monthly_summary(records):
        vals = [m["month"], m["sail_days"], m["port_days"], m["dock_days"],
                round(m["sail_fuel"], 1), round(m["port_fuel"], 1),
                round(m["total_fuel"], 1), round(m["refuel"], 1) or None]
        fills = [None, FILL_SAIL, FILL_PORT, FILL_DOCK,
                 FILL_SAIL, FILL_PORT, None, FILL_REFUEL if m["refuel"] else None]
        for col, (v, fill) in enumerate(zip(vals, fills), 1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = THIN
            if fill:
                c.fill = fill
        r += 1
    ws.freeze_panes = "A3"


def _sheet_daily(wb: Workbook, records: list[DailyRecord]) -> None:
    ws = wb.create_sheet("每日明細")
    _widths(ws, [11, 12, 7, 10, 9, 9, 9, 9, 9, 9, 9, 9, 10, 10])
    _title(ws, 1, "歷史逐日油耗明細（加油日黃底）", 14)
    _header(ws, 2, ["日期", "航次", "分類", "航行距離(NM)", "推進時數", "推進用油(KL)",
                    "作業時數", "作業用油(KL)", "靠港時數", "靠港用油(KL)",
                    "錨泊時數", "錨泊用油(KL)", "當日合計(KL)", "加油量(KL)"])
    r = 3
    for rec in records:
        vals = [f"{rec.day:%Y/%m/%d}", rec.voyage_no, rec.status, rec.distance,
                rec.prop_hrs, rec.prop_fuel, rec.oper_hrs, rec.oper_fuel,
                rec.port_hrs, rec.port_fuel, rec.anchor_hrs, rec.anchor_fuel,
                rec.total_fuel, rec.refuel or None]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = THIN
            if rec.refuel:
                c.fill = FILL_REFUEL
            elif col == 3:
                c.fill = STATUS_FILL.get(rec.status, FILL_PORT)
        r += 1
    ws.freeze_panes = "A3"


def _sheet_plan(wb: Workbook, plan: Plan2026) -> None:
    ws = wb.create_sheet("2026加油預估")
    _widths(ws, [11, 22, 16, 18, 11, 11, 9, 9, 13, 13, 13, 13])
    _title(ws, 1, "2026 各航次油耗預估與加油排程（依任務類型油耗率）", 12)
    _header(ws, 2, ["航次", "任務類型", "委託單位", "負責老師/備註", "出海日", "返港日",
                    "出海天數", "靠港天數", "出航油耗率(KL/日)", "預估用油(KL)",
                    "航段結束預估ROB(KL)", "加油（日期／量KL）"])
    r = 3
    for vp in plan.voyages:
        v = vp.voyage
        refuel = (f"{vp.refuel_day:%m/%d}／{vp.refuel_amount:.1f}"
                  if vp.refuel_day else "")
        vals = [v.voyage_no, v.vtype, v.client, v.pi,
                f"{v.start:%Y/%m/%d}", f"{v.end:%Y/%m/%d}",
                vp.sail_days, vp.port_days,
                vp.sail_rate if not v.is_dock else 0,
                vp.est_fuel, vp.end_rob, refuel]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=val)
            c.border = THIN
            if col == 2:
                c.fill = FILL_DOCK if v.is_dock else FILL_SAIL
            if col == 12 and refuel:
                c.fill = FILL_REFUEL
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="加油計畫彙總").font = FONT_TITLE; r += 1
    _header(ws, r, ["加油日", "觸發航次", "加油前ROB(KL)", "加油量(KL)",
                    "加油後ROB(KL)", "狀態", "", "", "", "", "", ""]); r += 1
    for ev in plan.refuels:
        vals = [f"{ev.day:%Y/%m/%d}", ev.after, ev.rob_before, ev.amount,
                ev.rob_after, "實際" if ev.is_actual else "預估"]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=val)
            c.border = THIN
            c.fill = FILL_REFUEL
        r += 1
    r += 1
    for alert in plan.alerts:
        ws.cell(row=r, column=1, value="⚠ " + alert).font = FONT_ALERT
        r += 1
    if not plan.alerts:
        ws.cell(row=r, column=1,
                value=f"✓ 全年預估 ROB 均高於警戒線 {WARNING_LEVEL_KL:.0f} KL")
    ws.freeze_panes = "A3"


def _sheet_rob(wb: Workbook, plan: Plan2026) -> None:
    ws = wb.create_sheet("2026逐日ROB")
    _widths(ws, [11, 4, 4, 11, 7, 12, 12, 11, 11, 11, 10])
    _title(ws, 1, ("2026 逐日油量｜灰底＝已有實際資料 黃底＝加油日｜"
                   "預估依任務類型油耗率自年初連續推算｜耗油誤差＝實際−預估"), 11)
    ws.cell(row=2, column=1, value=(
        f"年初 ROB {plan.initial_rob:.1f} KL｜油艙上限 {TANK_CAPACITY_KL:.0f} KL｜"
        f"警戒線 {WARNING_LEVEL_KL:.0f} KL"))
    _header(ws, 3, ["日期", "月", "日", "航次", "分類", "預估日耗油(KL)",
                    "實際日耗油(KL)", "預估ROB(KL)", "實際ROB(KL)",
                    "耗油誤差(KL)", "加油量(KL)"])
    r = 4
    for pd in plan.days:
        refuel = pd.refuel_actual if pd.refuel_actual is not None else (pd.refuel_est or None)
        vals = [f"{pd.day:%Y/%m/%d}", pd.day.month, pd.day.day,
                pd.voyage_no or None, pd.status, pd.est_fuel,
                pd.actual_fuel, pd.est_rob, pd.actual_rob,
                pd.fuel_error, refuel]
        has_actual = pd.actual_fuel is not None or pd.actual_rob is not None
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = THIN
            if refuel:
                c.fill = FILL_REFUEL
            elif has_actual:
                c.fill = FILL_ACTUAL
            elif col == 5:
                c.fill = STATUS_FILL.get(pd.status, FILL_PORT)
            if col == 8 and pd.below_warning:
                c.font = FONT_ALERT
        r += 1
    ws.freeze_panes = "A4"


def build_workbook(records: list[DailyRecord], rates: RateTable,
                   plan: Plan2026) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet_summary(wb, rates, records)
    _sheet_monthly(wb, records)
    _sheet_daily(wb, records)
    _sheet_plan(wb, plan)
    _sheet_rob(wb, plan)
    return wb
