"""歷史每日航務資料載入。

支援兩種來源：
1. SMF-07-05 TORI Daily log Abstract V2（原始格式，固定欄位序號）
   實際欄位對應（1-based）：
   - col 2  航次編號、col 5  日期（格式 "YYYYMMDD04:00"）
   - col 23 推進時數、col 24 作業時數、col 25 靠港時數
   - col 26 錨泊時數、col 27 航行距離
   - col 34 推進油耗、col 35 作業油耗、col 36 靠港油耗
   - col 37 錨泊油耗、col 38 當日合計油耗
   - col 39 加油量 MDO（填寫率極低，多數靠 ROB 差值判定）
   - col 43 ROB MDO
   資料行識別：col 1 = "Lengend" 或 "Legend"
   每個 sheet 覆蓋不同日期段，全部讀取合併後去重排序。

2. TORI_FuelConsumption_Summary 的「每日明細」工作表（已整理格式）

每日分類規則（固定順序）：
1. 當日總油耗 == 0 → 進塢
2. 推進時數 > 0 或 作業時數 > 0 或 距離 ≠ 0 或 推進油耗 > 0
   或 作業油耗 > 0 或 靠港油耗 == 0 → 出航
3. 其餘 → 靠港
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl

# ------------------------------------------------------------------
# SMF-07-05 V2 固定欄位（1-based；data 行 col1 = "Lengend"/"Legend"）
# ------------------------------------------------------------------
SMF_COL_SHIP = 1          # "Lengend"
SMF_COL_VOYAGE = 2        # 航次編號
SMF_COL_DATE = 5          # 日期，格式 "YYYYMMDD04:00"
SMF_COL_PROP_HRS = 23     # 推進時數
SMF_COL_OPER_HRS = 24     # 作業時數
SMF_COL_PORT_HRS = 25     # 靠港時數
SMF_COL_ANCHOR_HRS = 26   # 錨泊時數
SMF_COL_DISTANCE = 27     # 航行距離 (Miles)
SMF_COL_FUEL_PROP = 34    # 推進油耗 MDO/MGO
SMF_COL_FUEL_OPER = 35    # 作業油耗
SMF_COL_FUEL_PORT = 36    # 靠港油耗
SMF_COL_FUEL_ANCHOR = 37  # 錨泊油耗
SMF_COL_FUEL_TOTAL = 38   # 當日合計油耗
SMF_COL_REPLENISH = 39    # 加油量 MDO（填寫率低）
SMF_COL_ROB_MDO = 43      # ROB MDO

REFUEL_DELTA_KL = 10.0    # ROB 增加超過此值即判定為加油

STATUS_SAIL = "出航"
STATUS_PORT = "靠港"
STATUS_DOCK = "進塢"


@dataclass
class DailyRecord:
    day: date
    voyage_no: str = ""
    status: str = ""
    distance: float = 0.0
    prop_hrs: float = 0.0
    prop_fuel: float = 0.0
    oper_hrs: float = 0.0
    oper_fuel: float = 0.0
    port_hrs: float = 0.0
    port_fuel: float = 0.0
    anchor_hrs: float = 0.0
    anchor_fuel: float = 0.0
    total_fuel: float = 0.0
    refuel: float = 0.0
    rob: float | None = None


@dataclass
class Actual2026:
    """既有 2026 逐日表中的實際值（有資料的日期才有）。"""

    daily_fuel: dict[date, float] = field(default_factory=dict)
    rob: dict[date, float] = field(default_factory=dict)
    refuel: dict[date, float] = field(default_factory=dict)
    voyage: dict[date, str] = field(default_factory=dict)


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0


def _to_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    # SMF 格式：YYYYMMDD04:00 → 取前 8 碼
    if len(s) >= 8 and s[:8].isdigit():
        try:
            return datetime.strptime(s[:8], "%Y%m%d").date()
        except ValueError:
            pass
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def classify_day(prop_hrs, oper_hrs, distance, prop_fuel, oper_fuel,
                 port_fuel, total_fuel) -> str:
    if total_fuel == 0:
        return STATUS_DOCK
    if (prop_hrs > 0 or oper_hrs > 0 or distance != 0
            or prop_fuel > 0 or oper_fuel > 0 or port_fuel == 0):
        return STATUS_SAIL
    return STATUS_PORT


def detect_refuels(records: list[DailyRecord]) -> None:
    """以前後兩日 ROB 差值判定加油並回填加油量。

    加油量 = 當日 ROB − 前日 ROB + 當日油耗（即補回消耗後的淨增量）。
    直接填寫在記錄中的加油量（refuel > 0）優先，不覆蓋。
    """
    prev_rob: float | None = None
    for rec in records:
        if rec.rob is not None and prev_rob is not None:
            delta = rec.rob - prev_rob
            if delta > REFUEL_DELTA_KL and rec.refuel == 0:
                rec.refuel = round(delta + rec.total_fuel, 1)
        if rec.rob is not None:
            prev_rob = rec.rob


def _parse_smf_row(row: tuple) -> DailyRecord | None:
    """將 SMF V2 單一資料行解析為 DailyRecord；非資料行回傳 None。"""
    if len(row) < SMF_COL_ROB_MDO:
        return None
    ship = str(row[SMF_COL_SHIP - 1] or "").strip()
    if ship not in ("Lengend", "Legend"):
        return None
    day = _to_date(row[SMF_COL_DATE - 1])
    if day is None:
        return None
    rec = DailyRecord(
        day=day,
        voyage_no=str(row[SMF_COL_VOYAGE - 1] or "").strip(),
        prop_hrs=_num(row[SMF_COL_PROP_HRS - 1]),
        oper_hrs=_num(row[SMF_COL_OPER_HRS - 1]),
        port_hrs=_num(row[SMF_COL_PORT_HRS - 1]),
        anchor_hrs=_num(row[SMF_COL_ANCHOR_HRS - 1]),
        distance=_num(row[SMF_COL_DISTANCE - 1]),
        prop_fuel=_num(row[SMF_COL_FUEL_PROP - 1]),
        oper_fuel=_num(row[SMF_COL_FUEL_OPER - 1]),
        port_fuel=_num(row[SMF_COL_FUEL_PORT - 1]),
        anchor_fuel=_num(row[SMF_COL_FUEL_ANCHOR - 1]),
        total_fuel=_num(row[SMF_COL_FUEL_TOTAL - 1]),
        refuel=_num(row[SMF_COL_REPLENISH - 1]) if len(row) >= SMF_COL_REPLENISH else 0.0,
        rob=(None if row[SMF_COL_ROB_MDO - 1] in (None, "")
             else _num(row[SMF_COL_ROB_MDO - 1])),
    )
    rec.status = classify_day(rec.prop_hrs, rec.oper_hrs, rec.distance,
                              rec.prop_fuel, rec.oper_fuel,
                              rec.port_fuel, rec.total_fuel)
    return rec


def load_smf_daily_log(path: str, sheet: str | None = None) -> list[DailyRecord]:
    """讀取 SMF-07-05 V2 原始 Daily Log。

    預設讀取所有名稱含 "Daily" 的工作表並合併（不同 sheet 覆蓋不同日期段）。
    指定 sheet 名稱則只讀取該頁。同一日期若重複出現，保留最後一筆。
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = ([sheet] if sheet else
              [n for n in wb.sheetnames if "Daily" in n or "daily" in n])
    by_day: dict[date, DailyRecord] = {}
    for sname in sheets:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        for row in ws.iter_rows(values_only=True):
            rec = _parse_smf_row(row)
            if rec is not None:
                by_day[rec.day] = rec  # 後讀的覆蓋前讀的
    records = sorted(by_day.values(), key=lambda r: r.day)
    detect_refuels(records)
    return records


def merge_records(base: list[DailyRecord],
                  overlay: list[DailyRecord]) -> list[DailyRecord]:
    """合併兩個來源的記錄；overlay 中有的日期以 overlay 為準。"""
    by_day = {r.day: r for r in base}
    for r in overlay:
        by_day[r.day] = r
    return sorted(by_day.values(), key=lambda r: r.day)


def load_summary_daily(path: str, sheet: str = "每日明細") -> list[DailyRecord]:
    """讀取已整理的「每日明細」工作表。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    records: list[DailyRecord] = []
    for row in ws.iter_rows(values_only=True):
        day = _to_date(row[0])
        if day is None:
            continue
        records.append(DailyRecord(
            day=day,
            voyage_no=str(row[1] or ""),
            status=str(row[2] or ""),
            distance=_num(row[3]),
            prop_hrs=_num(row[4]), prop_fuel=_num(row[5]),
            oper_hrs=_num(row[6]), oper_fuel=_num(row[7]),
            port_hrs=_num(row[8]), port_fuel=_num(row[9]),
            anchor_hrs=_num(row[10]), anchor_fuel=_num(row[11]),
            total_fuel=_num(row[12]),
            refuel=_num(row[13]) if len(row) > 13 else 0.0,
        ))
    records.sort(key=lambda r: r.day)
    return records


def load_actual_2026(path: str, sheet: str = "2026逐日ROB") -> Actual2026:
    """從既有 2026 逐日表擷取實際日耗油、實際 ROB 與實際加油量。"""
    actual = Actual2026()
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        return actual
    ws = wb[sheet]
    for row in ws.iter_rows(values_only=True):
        day = _to_date(row[0])
        if day is None:
            continue
        has_actual = False
        if len(row) > 6 and row[6] is not None:
            actual.daily_fuel[day] = float(row[6])
            has_actual = True
        if len(row) > 8 and row[8] is not None:
            actual.rob[day] = float(row[8])
            has_actual = True
        # 加油量欄在未來日期放的是舊計畫的預估值，僅實際資料列才採計
        if has_actual and len(row) > 10 and row[10] is not None:
            actual.refuel[day] = float(row[10])
        if len(row) > 3 and row[3]:
            actual.voyage[day] = str(row[3])
    return actual
